"""
Excel export (Phase 6) — build a clean, employer-facing credit-spread workbook.

Consumes the deterministic ComputeResult (already computed from the HUMAN-REVIEWED
mapping) plus the ReviewState (for the audit trail) and the raw tables. openpyxl
only formats and writes — it never computes or alters a figure.

Styling conventions:
  * source/input values are BLUE, calculation outputs are BLACK
  * validation errors/warnings/info are colour-filled by severity
  * low-confidence mappings and human overrides are highlighted
  * every data sheet has a styled header, frozen header row, and an auto-filter
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.compute.engine import ComputeResult
from core.ingest.models import RawTable
from core.review import ReviewState
from core.schema import (
    CATEGORY_CATALOG,
    Severity,
    StandardizedCategory,
)

TOOL_NAME = "Credit Spread Builder"

# -- styles --
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
INPUT_FONT = Font(color="1F4E78")   # blue — source / input values
OUTPUT_FONT = Font(color="000000")  # black — calculation outputs
ERROR_FILL = PatternFill("solid", fgColor="FFC7CE")
ERROR_FONT = Font(color="9C0006", bold=True)
WARNING_FILL = PatternFill("solid", fgColor="FFEB9C")
WARNING_FONT = Font(color="9C6500")
INFO_FILL = PatternFill("solid", fgColor="DDEBF7")
LOWCONF_FILL = PatternFill("solid", fgColor="FCE4D6")
OVERRIDE_FILL = PatternFill("solid", fgColor="FFF2CC")

NUM_FMT = "#,##0.0;(#,##0.0)"
PCT_FMT = "0.0%"
RATIO_FMT = '0.00"x"'

_SEVERITY_STYLE = {
    Severity.ERROR: (ERROR_FILL, ERROR_FONT),
    Severity.WARNING: (WARNING_FILL, WARNING_FONT),
    Severity.INFO: (INFO_FILL, None),
}

# Catalog order for sorting spread rows.
_CAT_ORDER = {m.category: i for i, m in enumerate(CATEGORY_CATALOG)}


def _write_header(ws: Worksheet, headers: list[str], row: int = 1) -> None:
    for c, name in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _finish_table(ws: Worksheet, header_row: int, last_row: int, ncols: int,
                  widths: Optional[dict[int, int]] = None) -> None:
    last_col = get_column_letter(ncols)
    ws.freeze_panes = f"A{header_row + 1}"
    if last_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{last_col}{last_row}"
    for c in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 18)


# --- sheets -------------------------------------------------------------------
def _overview_sheet(ws, *, source_filename, income_table, balance_table,
                    generated_at, compute) -> None:
    ws.title = "Overview"
    ws["A1"] = TOOL_NAME
    ws["A1"].font = TITLE_FONT
    summary = compute.validation_summary()
    rows = [
        ("Tool", TOOL_NAME),
        ("Uploaded file", source_filename),
        ("Generated (UTC)", generated_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("Income Statement table", _table_desc(income_table)),
        ("Balance Sheet table", _table_desc(balance_table)),
        ("Reporting scale", (income_table.detected_scale or "unknown")),
        ("Reporting currency", (income_table.detected_currency or "unknown")),
        ("Validation", "PASS" if summary["passed"] else "FAIL"),
        ("Errors / Warnings / Info",
         f"{summary['n_errors']} / {summary['n_warnings']} / {summary['n_info']}"),
    ]
    r = 3
    for label, value in rows:
        ws.cell(row=r, column=1, value=label).font = BOLD
        ws.cell(row=r, column=2, value=value)
        r += 1

    notes = [
        "",
        "IMPORTANT NOTES:",
        "• For synthetic or public-company data only — never real borrower data.",
        "• AI (Claude) ONLY mapped raw line items to standardized categories.",
        "• All values, subtotals, ratios, and validation checks were computed "
        "deterministically in Python. Claude supplied no numbers.",
        "• Every figure traces to a source cell (see Mapping Audit Trail / Raw "
        "Extracted Rows).",
        "• This workbook REQUIRES human review before any decision or use.",
    ]
    for note in notes:
        cell = ws.cell(row=r, column=1, value=note)
        if note.endswith(":"):
            cell.font = BOLD
        r += 1
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60


def _table_desc(t: Optional[RawTable]) -> str:
    if t is None:
        return "(none)"
    return f"Table {t.table_index}: {t.heading or '(no heading)'}"


def _spread_sheet(ws, title: str, items, review: ReviewState) -> None:
    ws.title = title
    headers = ["standardized_category", "period", "raw_label", "source_value",
               "calculation_value", "confidence", "human_override", "source_reference"]
    _write_header(ws, headers)

    def _override(row_id: str) -> bool:
        try:
            return review.get(row_id).human_override
        except KeyError:
            return False

    ordered = sorted(
        [it for it in items if it.standardized_category != StandardizedCategory.UNMAPPED],
        key=lambda it: (_CAT_ORDER.get(it.standardized_category, 999), it.fiscal_year),
    )
    r = 2
    for it in ordered:
        overridden = _override(it.row_id)
        ws.cell(row=r, column=1, value=it.standardized_category.value)
        ws.cell(row=r, column=2, value=it.fiscal_year)
        ws.cell(row=r, column=3, value=it.raw_label)

        sv = ws.cell(row=r, column=4, value=it.source_value)
        sv.font = INPUT_FONT
        sv.number_format = NUM_FMT
        cv = ws.cell(row=r, column=5, value=it.calculation_value)
        cv.font = OUTPUT_FONT
        cv.number_format = NUM_FMT

        conf = ws.cell(row=r, column=6, value=it.confidence.value)
        if it.confidence.value == "low":
            conf.fill = LOWCONF_FILL
        ov = ws.cell(row=r, column=7, value="YES" if overridden else "")
        if overridden:
            ov.fill = OVERRIDE_FILL
            ov.font = BOLD
        ws.cell(row=r, column=8, value=it.source_location)
        r += 1

    _finish_table(ws, 1, r - 1, len(headers),
                  widths={1: 30, 3: 40, 8: 40})


def _ratios_sheet(ws, compute: ComputeResult) -> None:
    ws.title = "Ratios"
    periods = compute.periods
    headers = ["ratio"] + list(periods)
    _write_header(ws, headers)

    order = [
        ("current_ratio", RATIO_FMT),
        ("debt_to_equity", RATIO_FMT),
        ("debt_to_ebitda", RATIO_FMT),
        ("ebitda_margin", PCT_FMT),
        ("net_margin", PCT_FMT),
        ("revenue_growth", PCT_FMT),
        ("approximate_dscr", RATIO_FMT),
    ]
    by_period = compute.ratios_by_period()
    display = {}
    for rv in compute.ratios:
        display[rv.name] = rv.display_name

    r = 2
    for name, fmt in order:
        ws.cell(row=r, column=1, value=display.get(name, name)).font = BOLD
        for c, period in enumerate(periods, start=2):
            rv = by_period.get(period, {}).get(name)
            cell = ws.cell(row=r, column=c)
            if rv is not None and rv.status == "ok":
                cell.value = rv.value
                cell.number_format = fmt
                cell.font = OUTPUT_FONT
            else:
                cell.value = rv.message if rv else "Not calculated"
                cell.font = Font(italic=True, color="808080")
        r += 1

    ws.cell(row=r + 1, column=1,
            value="Approximate DSCR = EBITDA / (interest expense + current "
                  "portion of LTD). This is an approximation, NOT a bank-quality "
                  "or CFADS-based DSCR.").font = Font(italic=True)
    _finish_table(ws, 1, r - 1, len(headers), widths={1: 42})


def _validation_sheet(ws, compute: ComputeResult) -> None:
    ws.title = "Validation Checks"
    summary = compute.validation_summary()
    ws["A1"] = (f"Validation: {'PASS' if summary['passed'] else 'FAIL'}   "
                f"|   Errors: {summary['n_errors']}   "
                f"Warnings: {summary['n_warnings']}   Info: {summary['n_info']}")
    ws["A1"].font = BOLD

    header_row = 3
    headers = ["severity", "code", "message", "category", "period"]
    _write_header(ws, headers, row=header_row)

    order = {Severity.ERROR: 0, Severity.WARNING: 1, Severity.INFO: 2}
    findings = sorted(compute.findings, key=lambda f: order.get(f.severity, 3))

    r = header_row + 1
    for f in findings:
        fill, font = _SEVERITY_STYLE.get(f.severity, (None, None))
        sev = ws.cell(row=r, column=1, value=f.severity.value)
        if fill:
            sev.fill = fill
        if font:
            sev.font = font
        ws.cell(row=r, column=2, value=f.code.value)
        ws.cell(row=r, column=3, value=f.message)
        ws.cell(row=r, column=4, value=", ".join(f.related_categories))
        ws.cell(row=r, column=5, value=f.fiscal_year or "")
        r += 1

    if not findings:
        ws.cell(row=r, column=1, value="No validation findings.")
        r += 1

    _finish_table(ws, header_row, r - 1, len(headers),
                  widths={1: 12, 2: 26, 3: 70, 4: 28, 5: 10})


def _audit_sheet(ws, items, review: ReviewState) -> None:
    ws.title = "Mapping Audit Trail"
    headers = ["row_id", "statement_type", "raw_label", "original_category",
               "final_reviewed_category", "original_confidence", "human_override",
               "reviewer_note", "source_table_index", "source_row_index",
               "source_column_index", "source_location", "source_scale", "currency"]
    _write_header(ws, headers)

    r = 2
    for it in items:
        try:
            rm = review.get(it.row_id)
            original = rm.original_category.value
            orig_conf = rm.original_confidence.value
            override = rm.human_override
            note = rm.reviewer_note or ""
        except KeyError:
            original = it.standardized_category.value
            orig_conf = it.confidence.value
            override = False
            note = ""

        ws.cell(row=r, column=1, value=it.row_id)
        ws.cell(row=r, column=2, value=it.statement_type.value)
        ws.cell(row=r, column=3, value=it.raw_label)
        ws.cell(row=r, column=4, value=original)
        rev = ws.cell(row=r, column=5, value=it.standardized_category.value)
        conf = ws.cell(row=r, column=6, value=orig_conf)
        if orig_conf == "low":
            conf.fill = LOWCONF_FILL
        ov = ws.cell(row=r, column=7, value="YES" if override else "")
        if override:
            ov.fill = OVERRIDE_FILL
            ov.font = BOLD
            rev.fill = OVERRIDE_FILL
        ws.cell(row=r, column=8, value=note)
        ws.cell(row=r, column=9, value=it.source_table_index)
        ws.cell(row=r, column=10, value=it.source_row_index)
        ws.cell(row=r, column=11, value=it.source_column_index)
        ws.cell(row=r, column=12, value=it.source_location)
        ws.cell(row=r, column=13, value=it.source_scale.value)
        ws.cell(row=r, column=14, value=it.currency or "")
        r += 1

    _finish_table(ws, 1, r - 1, len(headers),
                  widths={3: 38, 4: 26, 5: 26, 8: 30, 12: 40})


def _raw_rows_sheet(ws, income_table: RawTable, balance_table: RawTable) -> None:
    ws.title = "Raw Extracted Rows"
    periods: list[str] = []
    for t in (income_table, balance_table):
        for p in t.period_labels:
            if p and p not in periods:
                periods.append(p)
    headers = ["statement", "table_index", "row_index", "kind", "raw_label"] + \
              [f"value[{p}]" for p in periods] + ["source_reference"]
    _write_header(ws, headers)

    r = 2
    for statement, t in (("income_statement", income_table),
                         ("balance_sheet", balance_table)):
        for prow in t.preview_rows():
            ws.cell(row=r, column=1, value=statement)
            ws.cell(row=r, column=2, value=t.table_index)
            ws.cell(row=r, column=3, value=prow.get("row_index"))
            ws.cell(row=r, column=4, value=prow.get("kind"))
            ws.cell(row=r, column=5, value=prow.get("label"))
            for c, p in enumerate(periods, start=6):
                cell = ws.cell(row=r, column=c, value=prow.get(p))
                cell.font = INPUT_FONT
                cell.number_format = NUM_FMT
            ws.cell(row=r, column=5 + len(periods) + 1, value=prow.get("source_ref"))
            r += 1

    _finish_table(ws, 1, r - 1, len(headers), widths={5: 42, 5 + len(periods) + 1: 24})


# --- entry points -------------------------------------------------------------
def build_workbook(
    *,
    source_filename: str,
    income_table: RawTable,
    balance_table: RawTable,
    compute: ComputeResult,
    review: ReviewState,
    generated_at: Optional[datetime] = None,
) -> Workbook:
    """Build the full credit-spread workbook from reviewed, computed data."""
    generated_at = generated_at or datetime.utcnow()
    wb = Workbook()

    _overview_sheet(wb.active, source_filename=source_filename,
                    income_table=income_table, balance_table=balance_table,
                    generated_at=generated_at, compute=compute)

    inc_items = [it for it in compute.items
                 if it.statement_type.value == "income_statement"]
    bal_items = [it for it in compute.items
                 if it.statement_type.value == "balance_sheet"]

    _spread_sheet(wb.create_sheet("Income Statement Spread"), "Income Statement Spread",
                  inc_items, review)
    _spread_sheet(wb.create_sheet("Balance Sheet Spread"), "Balance Sheet Spread",
                  bal_items, review)
    _ratios_sheet(wb.create_sheet("Ratios"), compute)
    _validation_sheet(wb.create_sheet("Validation Checks"), compute)
    _audit_sheet(wb.create_sheet("Mapping Audit Trail"), compute.items, review)
    _raw_rows_sheet(wb.create_sheet("Raw Extracted Rows"), income_table, balance_table)
    return wb


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_workbook_bytes(**kwargs) -> bytes:
    return workbook_to_bytes(build_workbook(**kwargs))
