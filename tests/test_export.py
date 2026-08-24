"""
Phase 6 — Excel export tests. Deterministic, no network, no LLM.

Builds the workbook from a reviewed + computed spread and reads it back to
verify structure and content.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from core.ingest import extract_document
from core.mapping import build_mapping_inputs, validate_and_bind
from core.compute import compute_spread
from core.review import ReviewState
from core.export import build_workbook_bytes
from core.schema import StandardizedCategory as C, ratios as ratios_spec

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "synthetic_filing.html"

EXPECTED_SHEETS = [
    "Overview", "Income Statement Spread", "Balance Sheet Spread", "Ratios",
    "Validation Checks", "Mapping Audit Trail", "Raw Extracted Rows",
]

_failures: list[str] = []

_MAP = {
    "Net sales": "revenue", "Cost of products sold": "cost_of_goods_sold",
    "Gross profit": "gross_profit",
    "Selling, general and administrative expenses": "selling_general_admin",
    "Operating income": "operating_income", "Interest expense": "interest_expense",
    "Other, net": "unmapped", "Income before income taxes": "income_before_taxes",
    "Provision for income taxes": "income_tax_expense", "Net income": "net_income",
    "Assets": "ignore", "Cash and cash equivalents": "cash_and_equivalents",
    "Trade receivables, net": "accounts_receivable", "Inventories": "inventory",
    "Total current assets": "total_current_assets", "Goodwill": "goodwill",
    "Other non-current assets": "other_noncurrent_assets", "Total assets": "total_assets",
    "Liabilities and Equity": "ignore", "Accounts payable": "accounts_payable",
    "Current portion of long-term debt": "current_portion_long_term_debt",
    "Total current liabilities": "total_current_liabilities",
    "Long-term debt": "long_term_debt",
    "Other non-current liabilities": "other_noncurrent_liabilities",
    "Total liabilities": "total_liabilities",
    "Total equity": "unmapped",  # left unmapped so we can override it
    "Total liabilities and equity": "total_liabilities_and_equity",
}


def check(name: str, cond: bool, extra: str = "") -> None:
    ok = bool(cond)
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" -- {extra}" if extra and not ok else ""))
    if not ok:
        _failures.append(name)


def _all_text(ws) -> list[str]:
    out = []
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if v is not None:
                out.append(str(v))
    return out


def _build():
    doc = extract_document("synthetic_filing.html", FIXTURE.read_bytes())
    inc, bal = doc.tables[0], doc.tables[1]
    inputs = build_mapping_inputs(inc, bal)
    mappings = [{
        "row_id": r.row_id,
        "standardized_category": _MAP.get(r.raw_label, "unmapped"),
        "confidence": "high" if _MAP.get(r.raw_label) not in (None, "unmapped") else "low",
        "rationale": "x",
    } for r in inputs]
    review = ReviewState.from_mapping_result(validate_and_bind({"mappings": mappings}, inputs))
    # Human override: map Total equity row to total_equity.
    eq_rid = next(r.row_id for r in review.rows() if r.raw_label == "Total equity")
    review.set_category(eq_rid, C.TOTAL_EQUITY, note="analyst confirmed equity")
    compute = compute_spread(inc, bal, review.to_decisions_by_id(),
                             notes_by_id=review.notes_by_id())
    data = build_workbook_bytes(source_filename="synthetic_filing.html",
                                income_table=inc, balance_table=bal,
                                compute=compute, review=review)
    return load_workbook(io.BytesIO(data))


def test_workbook_and_sheets() -> None:
    print("workbook + sheets:")
    wb = _build()
    check("workbook created (has sheets)", len(wb.sheetnames) > 0)
    check("all expected sheets exist", wb.sheetnames == EXPECTED_SHEETS,
          str(wb.sheetnames))


def test_spread_sheets_have_both_values() -> None:
    print("spread sheets show source_value AND calculation_value:")
    wb = _build()
    for sheet in ("Income Statement Spread", "Balance Sheet Spread"):
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        check(f"{sheet} has source_value col", "source_value" in header)
        check(f"{sheet} has calculation_value col", "calculation_value" in header)
        # at least one data row with both numeric values populated
        svi = header.index("source_value")
        cvi = header.index("calculation_value")
        has_row = any(
            row[svi] is not None and row[cvi] is not None
            for row in ws.iter_rows(min_row=2, values_only=True)
        )
        check(f"{sheet} has a populated data row", has_row)


def test_audit_trail_original_and_reviewed() -> None:
    print("audit trail shows original + reviewed + override:")
    wb = _build()
    ws = wb["Mapping Audit Trail"]
    header = [c.value for c in ws[1]]
    check("has original_category", "original_category" in header)
    check("has final_reviewed_category", "final_reviewed_category" in header)
    oi = header.index("original_category")
    ri = header.index("final_reviewed_category")
    hi = header.index("human_override")
    li = header.index("raw_label")
    override_row = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[li] == "Total equity":
            override_row = row
            break
    check("Total equity row present in audit", override_row is not None)
    if override_row:
        check("original category is unmapped", override_row[oi] == "unmapped")
        check("reviewed category is total_equity", override_row[ri] == "total_equity")
        check("human_override marked YES", override_row[hi] == "YES")


def test_ratios_not_calculated_text() -> None:
    print("ratios sheet shows not-calculated text for DSCR:")
    wb = _build()
    ws = wb["Ratios"]
    text = " ".join(_all_text(ws))
    # Fixture has no D&A row, so EBITDA and DSCR are not calculable.
    check("DSCR not-calculated message present",
          ratios_spec.NOT_CALCULATED in text, ratios_spec.NOT_CALCULATED)
    check("DSCR label present", "Approximate DSCR" in text)


def test_validation_findings_written() -> None:
    print("validation findings written:")
    wb = _build()
    ws = wb["Validation Checks"]
    header_texts = _all_text(ws)
    check("has a PASS/FAIL summary", any("Validation:" in t for t in header_texts))
    # header row is row 3
    header = [c.value for c in ws[3]]
    check("severity column present", "severity" in header)
    check("code column present", "code" in header)
    # at least one finding row (unmapped Other,net / low confidence / footing / ratio)
    body = list(ws.iter_rows(min_row=4, values_only=True))
    check("at least one finding row", any(r[0] for r in body))


def test_overview_notes() -> None:
    print("overview notes:")
    wb = _build()
    text = " ".join(_all_text(wb["Overview"]))
    check("mentions synthetic/public-company only", "synthetic" in text.lower())
    check("states AI only mapped categories", "mapped" in text.lower() and "Claude" in text)
    check("states human review required", "human review" in text.lower())
    check("shows uploaded filename", "synthetic_filing.html" in text)


def main() -> int:
    test_workbook_and_sheets()
    test_spread_sheets_have_both_values()
    test_audit_trail_original_and_reviewed()
    test_ratios_not_calculated_text()
    test_validation_findings_written()
    test_overview_notes()
    print()
    if _failures:
        print(f"FAILED ({len(_failures)}): " + ", ".join(_failures))
        return 1
    print("All Phase 6 export checks passed.")
    return 0


def test_all():
    assert main() == 0


if __name__ == "__main__":
    raise SystemExit(main())
